"""生成《幽林 · 网页开发学习记录与未来目标》DOCX。

内容来源：D:\\my-site 的 11 课学习记录、提交历史、进度表。
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_CN = "微软雅黑"
FONT_EN = "Segoe UI"
ACCENT = RGBColor(0x2F, 0x6B, 0xFF)
GREY = RGBColor(0x66, 0x66, 0x66)
DARK = RGBColor(0x22, 0x22, 0x22)


def set_run_font(run, size=None, bold=None, color=None, name_en=FONT_EN, name_cn=FONT_CN):
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_base(document):
    normal = document.styles["Normal"]
    normal.font.name = FONT_EN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(6)

    for lvl, size in ((1, 16), (2, 12.5), (3, 11)):
        st = document.styles[f"Heading {lvl}"]
        st.font.name = FONT_EN
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = DARK
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        st.paragraph_format.space_before = Pt(16 if lvl == 1 else 12)
        st.paragraph_format.space_after = Pt(8)


def para(document, text="", size=10.5, bold=False, color=None, align=None,
         space_after=None, indent=None):
    p = document.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return p


def bullet(document, text, bold_prefix="", size=10.5):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r0 = p.add_run(bold_prefix)
        set_run_font(r0, size=size, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def num_item(document, num, text, size=10.5):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.space_after = Pt(4)
    r0 = p.add_run(f"{num}. ")
    set_run_font(r0, size=size, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def hrule(document):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D8DCE3")
    borders.append(bottom)
    pPr.append(borders)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def make_table(document, headers, rows, widths, font_size=9.5, header_fill="EEF2F8"):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 固定表格布局，否则 LibreOffice / Word 会按内容重新分配列宽
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)

    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        shade(hdr[i], header_fill)
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True, color=DARK)

    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            r = p.add_run(str(val))
            set_run_font(r, size=font_size)

    for row in table.rows:
        # 禁止单行被拆到两页
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)

    # 首行作为表头，跨页时自动重复
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table


def build(path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    style_base(doc)

    # ---------------- 封面标题区 ----------------
    p = para(doc, "网页开发学习记录与未来目标", size=22, bold=True, color=ACCENT, space_after=4)
    para(doc, "幽林 · 计算机科学与技术 大一", size=11, color=GREY, space_after=2)
    para(doc, "整理时间：2026 年 10 月 20 日", size=9.5, color=GREY, space_after=2)
    hrule(doc)

    # ---------------- 一、个人现状 ----------------
    doc.add_heading("一、我现在的状态", level=1)
    para(doc, "从零基础开始，用 11 节课、四十余次提交，把一个纯手写的个人网站做到了公网可访问。"
              "以下是截至今天的事实，不是计划。")

    bullet(doc, "计算机科学与技术专业，大一。当前重点是专业基础、编程练习和英语。", "身份：")
    bullet(doc, "C++（专业基础课在读）、Python（在学爬虫第 08 节「网页数据的种类划分」）。"
                "网页开发在 10 月前完全没碰过。", "技术底子：")
    bullet(doc, "作曲 / MIDI 制作、健身。音乐是长期爱好，不是副业。", "长期投入：")
    bullet(doc, "项目驱动。每学一个知识点就落到网站上，看得见才算学会。", "学习方式：")
    bullet(doc, "重视实际能力、作品产出和时间自主权；不做对能力无增量的事。", "取舍原则：")

    # ---------------- 二、学习路径 ----------------
    doc.add_heading("二、学习路径：11 课，从「网页是什么」到「用数据生成页面」", level=1)
    para(doc, "每一课都有教材、动手任务和验收标准。下表按时间顺序记录了学到什么、做出什么。",
         color=GREY)

    rows = [
        ("1", "网页是怎么工作的", "客户端-服务器模型、file:// 与 http:// 的区别、F12 Network 面板", "本地服务器跑起来"),
        ("2", "HTML 骨架", "标签配对与嵌套、属性、head / body 的分工", "第一张个人页面"),
        ("3", "CSS 入门", "外部样式表、标签 / class / 伪类选择器、优先级意识", "三页样式同改"),
        ("4", "盒子模型 + Flexbox", "content / padding / border / margin、box-sizing、flex 主轴与交叉轴、gap", "卡片横向布局"),
        ("5", "多页面网站", "语义标签、一份 CSS 多页共用、相对路径 ../、视口单位", "四页完整站点"),
        ("6", "JavaScript 入门", "let / const、函数、DOM 查询、事件监听、classList", "暗色模式、点击计数"),
        ("7", "把网站变成真的", "audio 标签、图片资源、文件命名规范、真内容填充", "作品 / 音乐 / 笔记三页"),
        ("8", "Git 版本管理", "工作区-暂存区-仓库、commit、diff、restore、.gitignore", "28 个文件的版本历史"),
        ("9", "上线 GitHub Pages", "远程仓库、SSH 密钥认证、Pages 构建、HTTP 状态验证", "公网可访问的正式站点"),
        ("10", "响应式 + 分享卡片", "viewport、clamp() 流体字号、媒体查询、og 协议", "手机端不再横向溢出"),
        ("11", "用数据生成页面", "数组与对象、map / join、fetch 与异步、innerHTML 与 XSS", "曲目列表数据驱动"),
    ]
    make_table(doc, ["课次", "主题", "关键知识点", "产出"],
               rows, widths=[1.35, 3.3, 8.0, 3.95])

    para(doc, "说明：第 11 课是当前进度，示例与练习已交付，任务在推进中。", size=9, color=GREY)

    # ---------------- 三、作品 ----------------
    doc.add_heading("三、作品：一个纯手写的个人网站", level=2)
    para(doc, "没有框架、没有模板、没有构建工具。四个页面加一个音频播放器，全部手写。")
    bullet(doc, "https://small-handsome001.github.io/mysite/", "线上地址：")
    bullet(doc, "https://github.com/small-handsome001/mysite", "源码仓库：")
    bullet(doc, "HTML / CSS / 原生 JavaScript / Git / GitHub Pages", "技术栈：")
    bullet(doc, "首页、项目、音乐、笔记，另有 3 篇独立笔记文章页", "页面结构：")
    bullet(doc, "头像原图 787 KB，压缩为 240×240 后 15 KB，加载从 27 秒降到 1 秒", "一次真实的性能优化：")
    bullet(doc, "加一首曲子、写一篇笔记，只改一处文件，无需动结构代码（第 11 课目标）", "内容与结构分离：")

    # ---------------- 四、踩过的坑 ----------------
    doc.add_heading("四、踩过的坑：这张表比成绩单有用", level=1)
    para(doc, "每一个都是我自己犯的，每一个都写成了笔记或记得住的经验。", color=GREY)

    pit_rows = [
        ("改了代码页面没变", "忘了 Ctrl+S 保存", "先看是否已保存，再怀疑代码"),
        ("页面空白 / 图片跑到最上面", "<img> 写进了 <head>", "head 只放给浏览器看的信息"),
        ("每页顶部多一个反斜杠", "标签行末多了 \\ 字符", "HTML 不报错，多余字符会被当正文渲染"),
        ("音频一直点不动", "文件名成了 track-01.mp3.mp3", "Windows 隐藏扩展名造成的双后缀"),
        ("卡片被一张图撑爆", "图片没有 class，也未限制宽度", "全局 img { max-width: 100% }"),
        ("width: 120px 不生效", "img.avatar 权重高于 .avatar", "F12 里被划掉的规则就是输掉的那条"),
        ("CSS 后半段整体失效", "注释少了闭合 */", "数一数 /* 与 */ 是否配对"),
        ("SSH 说接受了钥匙又拒绝", "私钥密码被写成两个引号字符", "Server accepts key + Permission denied = 私钥打不开"),
        ("推送成功但线上没变", "Pages 构建在排队（Queued）", "先看 Actions 状态，不要在没证据时重复推送"),
        ("以为没部署成功", "curl 读到了 CDN 缓存", "看 Last-Modified 或加 no-cache 请求头"),
    ]
    make_table(doc, ["现象", "真正的原因", "留下的教训"],
               pit_rows, widths=[4.6, 5.4, 6.6])

    # ---------------- 五、未来目标 ----------------
    doc.add_heading("五、未来目标", level=1)

    doc.add_heading("近期（1—3 个月）", level=2)
    bullet(doc, "完成第 11 课任务：音乐页改为数据驱动，加曲子只改 JSON 一处。")
    bullet(doc, "继续第 12—14 课：JS 再深一层（数组方法、localStorage、fetch 真实数据）、"
                "把重复的页头页脚用 JS 或模板消除、站点性能与可访问性收尾。")
    bullet(doc, "内容从「占位」变「真实」：3 首原创曲目、3 篇真正踩过坑的笔记、"
                "1 个 C++ 小项目并配上 README。")
    bullet(doc, "Python 侧完成爬虫基础和进阶剩余节点中真正有用的部分（异步协程、抓包工具），"
                "跳过与自己目标无关的重型框架。")

    doc.add_heading("中期（大二前）", level=2)
    bullet(doc, "把网站做成可以拿给别人看的作品集：不是「我在学」，而是「这是我做的」。")
    bullet(doc, "写出一个对自己真正有用的 Python 工具（例如 MIDI 批量处理），"
                "标准是：自己愿意天天用。")
    bullet(doc, "英语达到能读懂官方文档、能看懂英文技术视频的水平；"
                "这是获取一手信息的门槛，不是考试任务。")

    doc.add_heading("长期（毕业前）", level=2)
    bullet(doc, "具备独立完成一个小型全栈应用的能力：前端界面、数据存储、部署上线，一个人能跑通。")
    bullet(doc, "积累 3 个以上有代码、有 README、能跑起来的项目，而不是 10 个半途而废的作业。")
    bullet(doc, "保持作曲和健身两条线不中断——它们不是「业余爱好」，是长期产出的一部分。")

    doc.add_heading("明确不做的事", level=2)
    para(doc, "写下来是为了避免被默认路径带走：")
    bullet(doc, "不为简历参加没有实际产出的活动；")
    bullet(doc, "不走考公、事业单位路线，不追求行政职务；")
    bullet(doc, "不为了「看起来在学」而囤课，学完一节就用一节。")

    # ---------------- 六、方法 ----------------
    doc.add_heading("六、这套方法里真正起作用的部分", level=1)
    num_item(doc, 1, "先用最笨的工具（原生 HTML/CSS/JS），把反馈闭环压到最短：改一行、刷新、立刻看到。")
    num_item(doc, 2, "每课都有「能验证的产出」，不是「看完了」。")
    num_item(doc, 3, "所有错误都记下来。同样的错不犯第二次，比多学一个知识点更值钱。")
    num_item(doc, 4, "遇到异常先看证据（F12、Actions、日志），再看代码。猜是最慢的调试方式。")
    num_item(doc, 5, "内容用数据描述，结构用代码生成。重复劳动交给程序，这条不只适用于网页。")

    hrule(doc)
    para(doc, "本文档由学习记录自动整理，内容与仓库提交历史一致。", size=9, color=GREY)
    para(doc, "仓库：github.com/small-handsome001/mysite　·　站点：small-handsome001.github.io/mysite/",
         size=9, color=GREY)

    doc.save(path)
    print("saved", path)


if __name__ == "__main__":
    build(r"D:\my-site\幽林-网页开发学习记录与未来目标.docx")



